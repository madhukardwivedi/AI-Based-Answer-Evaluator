from __future__ import annotations

import csv
from io import BytesIO, StringIO
from pathlib import Path
from uuid import UUID

from fastapi import HTTPException, UploadFile
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.models.student import Student
from app.models.user import User
from app.repositories.student_repository import StudentRepository
from app.core.config import get_settings
from app.schemas.student import StudentCreate, StudentImportError, StudentImportResult, StudentUpdate
from app.services.exam_service import ExamService


class StudentService:
    def __init__(self, students: StudentRepository | None = None, exams: ExamService | None = None) -> None:
        self.students = students or StudentRepository()
        self.exams = exams or ExamService()

    def list(self, db: Session, exam_id: UUID, user: User) -> list[Student]:
        self.exams.get(db, exam_id, user)
        return self.students.list_for_exam(db, exam_id)

    def get(self, db: Session, student_id: UUID, user: User) -> Student:
        student = self.students.get(db, student_id)
        if student is None:
            raise HTTPException(status_code=404, detail="Student not found")
        self.exams.get(db, student.exam_id, user)
        return student

    def create(self, db: Session, exam_id: UUID, data: StudentCreate, user: User) -> Student:
        self.exams.get(db, exam_id, user)
        return self._save(db, Student(exam_id=exam_id, **data.model_dump()))

    def update(self, db: Session, student_id: UUID, data: StudentUpdate, user: User) -> Student:
        student = self.get(db, student_id, user)
        for field, value in data.model_dump().items():
            setattr(student, field, value)
        return self._save(db, student)

    def delete(self, db: Session, student_id: UUID, user: User) -> None:
        self.students.delete(db, self.get(db, student_id, user))

    def import_roster(self, db: Session, exam_id: UUID, upload: UploadFile, user: User) -> StudentImportResult:
        self.exams.get(db, exam_id, user)
        filename = Path(upload.filename or "").name
        extension = Path(filename).suffix.lower()
        if extension not in {".csv", ".xlsx"}:
            raise HTTPException(status_code=422, detail="Only CSV and XLSX student roster files are supported")

        maximum_size = get_settings().max_student_import_size_bytes
        content = upload.file.read(maximum_size + 1)
        if not content:
            raise HTTPException(status_code=422, detail="Student roster file must not be empty")
        if len(content) > maximum_size:
            raise HTTPException(status_code=422, detail="Student roster file exceeds the configured import size limit")

        rows = self._parse_roster(content, extension)
        existing_roll_numbers, existing_external_ids = self.students.existing_identifiers(db, exam_id)
        seen_roll_numbers: set[str] = set()
        seen_external_ids: set[str] = set()
        errors: list[StudentImportError] = []
        students_to_create: list[Student] = []

        for row_number, row in rows:
            name = row.get("name", "").strip()
            roll_number = row.get("roll_number") or None
            external_id = row.get("external_id") or None
            row_errors: list[tuple[str, str]] = []
            if not name:
                row_errors.append(("name", "name is required"))
            if roll_number:
                if roll_number in existing_roll_numbers:
                    row_errors.append(("roll_number", f"roll_number {roll_number} already exists in this exam"))
                elif roll_number in seen_roll_numbers:
                    row_errors.append(("roll_number", f"roll_number {roll_number} is duplicated in this file"))
            if external_id:
                if external_id in existing_external_ids:
                    row_errors.append(("external_id", f"external_id {external_id} already exists in this exam"))
                elif external_id in seen_external_ids:
                    row_errors.append(("external_id", f"external_id {external_id} is duplicated in this file"))
            if row_errors:
                errors.extend(StudentImportError(row=row_number, field=field, message=message) for field, message in row_errors)
                if roll_number:
                    seen_roll_numbers.add(roll_number)
                if external_id:
                    seen_external_ids.add(external_id)
                continue
            students_to_create.append(Student(exam_id=exam_id, name=name, roll_number=roll_number, external_id=external_id))
            if roll_number:
                seen_roll_numbers.add(roll_number)
            if external_id:
                seen_external_ids.add(external_id)

        try:
            db.add_all(students_to_create)
            db.commit()
        except IntegrityError as error:
            db.rollback()
            # This can only happen if another import/manual creation raced this request.
            raise HTTPException(status_code=409, detail="A student identifier was added concurrently; please retry the import") from error

        return StudentImportResult(
            total_rows=len(rows),
            imported_count=len(students_to_create),
            skipped_count=len(rows) - len(students_to_create),
            errors=errors,
        )

    @staticmethod
    def _parse_roster(content: bytes, extension: str) -> list[tuple[int, dict[str, str]]]:
        if extension == ".csv":
            try:
                text = content.decode("utf-8-sig")
            except UnicodeDecodeError as error:
                raise HTTPException(status_code=422, detail="CSV file must be UTF-8 encoded") from error
            try:
                return StudentService._normalize_rows(csv.reader(StringIO(text), strict=True))
            except csv.Error as error:
                raise HTTPException(status_code=422, detail="CSV file could not be parsed") from error

        try:
            from openpyxl import load_workbook
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            try:
                worksheet = workbook.worksheets[0]
                return StudentService._normalize_rows(worksheet.iter_rows(values_only=True))
            finally:
                workbook.close()
        except HTTPException:
            raise
        except Exception as error:
            raise HTTPException(status_code=422, detail="XLSX file could not be parsed") from error

    @staticmethod
    def _normalize_rows(raw_rows: object) -> list[tuple[int, dict[str, str]]]:
        iterator = iter(raw_rows)  # type: ignore[arg-type]
        header_row_number = 0
        headers: list[str] | None = None
        for header_row_number, row in enumerate(iterator, start=1):
            values = ["" if value is None else str(value).strip() for value in row]
            if any(values):
                headers = values
                break
        if headers is None:
            raise HTTPException(status_code=422, detail="Student roster file does not contain headers")

        aliases = {
            "name": "name",
            "roll number": "roll_number",
            "roll_number": "roll_number",
            "roll no": "roll_number",
            "external id": "external_id",
            "external_id": "external_id",
        }
        canonical_headers = [aliases.get(header.casefold()) for header in headers]
        if "name" not in canonical_headers:
            raise HTTPException(status_code=422, detail="Student roster must include a name column")
        result: list[tuple[int, dict[str, str]]] = []
        for row_number, row in enumerate(iterator, start=header_row_number + 1):
            values = ["" if value is None else str(value).strip() for value in row]
            if not any(values):
                continue
            normalized: dict[str, str] = {}
            for index, canonical in enumerate(canonical_headers):
                if canonical is not None and canonical not in normalized:
                    normalized[canonical] = values[index] if index < len(values) else ""
            result.append((row_number, normalized))
        return result

    def _save(self, db: Session, student: Student) -> Student:
        try:
            return self.students.save(db, student)
        except IntegrityError as error:
            db.rollback()
            raise HTTPException(status_code=409, detail="Roll number already exists for this exam") from error
